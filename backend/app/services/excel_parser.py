"""Excel parser and validation service for sanctioned inventory baseline files."""

import io
from typing import Optional, List, Dict, Set
import openpyxl
from backend.app.schemas.inventory import (
    InventoryParsedItem,
    InventoryRowError,
    InventoryUploadPreviewResponse,
)

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB limit


def parse_and_validate_inventory_excel(
    file_bytes: bytes,
    filename: str,
) -> InventoryUploadPreviewResponse:
    """Parses and validates a sanctioned inventory Excel spreadsheet (.xlsx).
    
    Enforces required columns, data types, non-negative integers, and detects duplicates.
    Does NOT write to the database (preview only).
    """
    errors: List[InventoryRowError] = []
    parsed_items: List[InventoryParsedItem] = []

    # 1. Filename extension check
    if not filename.lower().endswith(".xlsx"):
        return InventoryUploadPreviewResponse(
            is_valid=False,
            filename=filename,
            total_rows=0,
            valid_rows_count=0,
            error_count=1,
            errors=[
                InventoryRowError(
                    row=0,
                    field="filename",
                    message="Unsupported file format. Only .xlsx Excel files are supported.",
                )
            ],
        )

    # 2. File size check
    if len(file_bytes) == 0:
        return InventoryUploadPreviewResponse(
            is_valid=False,
            filename=filename,
            total_rows=0,
            valid_rows_count=0,
            error_count=1,
            errors=[
                InventoryRowError(
                    row=0,
                    field="file",
                    message="Uploaded Excel file is empty (0 bytes).",
                )
            ],
        )

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        return InventoryUploadPreviewResponse(
            is_valid=False,
            filename=filename,
            total_rows=0,
            valid_rows_count=0,
            error_count=1,
            errors=[
                InventoryRowError(
                    row=0,
                    field="file",
                    message=f"File exceeds maximum allowed size of {MAX_FILE_SIZE_BYTES // (1024 * 1024)}MB.",
                )
            ],
        )

    # 3. Open workbook
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        return InventoryUploadPreviewResponse(
            is_valid=False,
            filename=filename,
            total_rows=0,
            valid_rows_count=0,
            error_count=1,
            errors=[
                InventoryRowError(
                    row=0,
                    field="file",
                    message=f"Invalid or corrupted Excel file: {str(exc)}",
                )
            ],
        )

    sheet = workbook.active
    if sheet is None or sheet.max_row is None or sheet.max_row < 1:
        return InventoryUploadPreviewResponse(
            is_valid=False,
            filename=filename,
            total_rows=0,
            valid_rows_count=0,
            error_count=1,
            errors=[
                InventoryRowError(
                    row=0,
                    field="worksheet",
                    message="Worksheet is empty and contains no data rows.",
                )
            ],
        )

    # 4. Find header row
    header_row_idx: Optional[int] = None
    item_col_idx: Optional[int] = None
    qty_col_idx: Optional[int] = None
    cat_col_idx: Optional[int] = None

    # Supported header aliases
    ITEM_ALIASES = {"item name", "item", "item_name", "asset", "equipment", "asset name", "name"}
    QTY_ALIASES = {"expected quantity", "quantity", "qty", "expected", "expected_qty", "sanctioned quantity", "sanctioned"}
    CAT_ALIASES = {"category", "type", "item category", "department"}

    found_headers: List[str] = []

    for r in range(1, min(sheet.max_row + 1, 15)):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
        cleaned_vals = [str(v).strip().lower() for v in row_vals if v is not None and str(v).strip()]
        
        # Check if row looks like a header containing at least item or qty
        has_item = any(v in ITEM_ALIASES for v in cleaned_vals)
        has_qty = any(v in QTY_ALIASES for v in cleaned_vals)

        if has_item or has_qty:
            header_row_idx = r
            for c in range(1, sheet.max_column + 1):
                val = sheet.cell(row=r, column=c).value
                if val is not None:
                    norm = str(val).strip().lower().replace("_", " ")
                    found_headers.append(str(val).strip())
                    if norm in ITEM_ALIASES and item_col_idx is None:
                        item_col_idx = c
                    elif norm in QTY_ALIASES and qty_col_idx is None:
                        qty_col_idx = c
                    elif norm in CAT_ALIASES and cat_col_idx is None:
                        cat_col_idx = c
            break

    if header_row_idx is None or item_col_idx is None or qty_col_idx is None:
        missing = []
        if item_col_idx is None:
            missing.append("'Item Name'")
        if qty_col_idx is None:
            missing.append("'Expected Quantity'")
        return InventoryUploadPreviewResponse(
            is_valid=False,
            filename=filename,
            total_rows=0,
            valid_rows_count=0,
            error_count=1,
            errors=[
                InventoryRowError(
                    row=header_row_idx or 1,
                    field="header",
                    message=f"Missing required column(s): {' and '.join(missing)}. Detected columns: {', '.join(found_headers) if found_headers else 'None'}.",
                )
            ],
        )

    # 5. Parse and validate data rows
    seen_items: Set[str] = set()
    total_data_rows = 0

    for r in range(header_row_idx + 1, sheet.max_row + 1):
        # Read cells
        item_val = sheet.cell(row=r, column=item_col_idx).value
        qty_val = sheet.cell(row=r, column=qty_col_idx).value
        cat_val = sheet.cell(row=r, column=cat_col_idx).value if cat_col_idx else None

        # Check if row is completely empty
        if item_val is None and qty_val is None and cat_val is None:
            continue

        total_data_rows += 1
        row_has_error = False

        # Validate Item Name
        if item_val is None or not str(item_val).strip():
            errors.append(
                InventoryRowError(
                    row=r,
                    field="Item Name",
                    message=f"Row {r}: Item Name is missing or empty.",
                )
            )
            row_has_error = True
            cleaned_item_name = ""
        else:
            cleaned_item_name = str(item_val).strip()
            if len(cleaned_item_name) > 100:
                errors.append(
                    InventoryRowError(
                        row=r,
                        field="Item Name",
                        message=f"Row {r}: Item Name exceeds maximum 100 characters.",
                    )
                )
                row_has_error = True

            # Duplicate check within file
            norm_item = cleaned_item_name.lower()
            if norm_item in seen_items:
                errors.append(
                    InventoryRowError(
                        row=r,
                        field="Item Name",
                        message=f"Row {r}: Duplicate item '{cleaned_item_name}' detected. Please consolidate duplicate entries in the Excel file.",
                    )
                )
                row_has_error = True
            else:
                seen_items.add(norm_item)

        # Validate Expected Quantity
        if qty_val is None or str(qty_val).strip() == "":
            errors.append(
                InventoryRowError(
                    row=r,
                    field="Expected Quantity",
                    message=f"Row {r}: Expected Quantity is missing.",
                )
            )
            row_has_error = True
            parsed_qty = 0
        else:
            try:
                # Handle floats like 20.0 from Excel numeric cells
                f_val = float(qty_val)
                if not f_val.is_integer() or isinstance(qty_val, bool):
                    errors.append(
                        InventoryRowError(
                            row=r,
                            field="Expected Quantity",
                            message=f"Row {r}: Expected Quantity must be a whole integer (got '{qty_val}').",
                        )
                    )
                    row_has_error = True
                    parsed_qty = 0
                elif f_val < 0:
                    errors.append(
                        InventoryRowError(
                            row=r,
                            field="Expected Quantity",
                            message=f"Row {r}: Expected Quantity must be a non-negative integer (got '{qty_val}').",
                        )
                    )
                    row_has_error = True
                    parsed_qty = 0
                else:
                    parsed_qty = int(f_val)
            except (ValueError, TypeError):
                errors.append(
                    InventoryRowError(
                        row=r,
                        field="Expected Quantity",
                        message=f"Row {r}: Expected Quantity must be numeric (got '{qty_val}').",
                    )
                )
                row_has_error = True
                parsed_qty = 0

        # Category
        cleaned_cat = str(cat_val).strip() if cat_val is not None and str(cat_val).strip() else None
        if cleaned_cat and len(cleaned_cat) > 100:
            cleaned_cat = cleaned_cat[:100]

        if not row_has_error and cleaned_item_name:
            parsed_items.append(
                InventoryParsedItem(
                    row=r,
                    item_name=cleaned_item_name,
                    category=cleaned_cat,
                    expected_quantity=parsed_qty,
                )
            )

    if total_data_rows == 0:
        errors.append(
            InventoryRowError(
                row=header_row_idx + 1 if header_row_idx else 1,
                field="rows",
                message="Excel sheet contains header but has 0 data rows.",
            )
        )

    is_valid = len(errors) == 0 and len(parsed_items) > 0

    return InventoryUploadPreviewResponse(
        is_valid=is_valid,
        filename=filename,
        total_rows=total_data_rows,
        valid_rows_count=len(parsed_items),
        error_count=len(errors),
        items=parsed_items,
        errors=errors,
    )
